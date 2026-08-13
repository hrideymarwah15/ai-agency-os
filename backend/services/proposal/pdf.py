"""Local proposal document rendering with optional dependencies and bounded concurrency."""

from __future__ import annotations

import asyncio
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

_PRINT_STYLE = "<style>@page{size:A4;margin:18mm}table,.proposal-section{break-inside:avoid}@media print{body{margin:0}}</style>"
_PDF_SEMAPHORE = asyncio.Semaphore(2)


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[tuple[str, str]] = []
        self._tag = "p"
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"h1", "h2", "h3", "p", "li"}:
            self._tag = tag
            self._text = []

    def handle_data(self, data: str) -> None:
        self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == self._tag and (text := " ".join("".join(self._text).split())):
            self.blocks.append((tag, text))
            self._text = []


class PDFRenderer:
    """Render HTML, PDF, and optional DOCX without requiring an external service."""

    def render_html(self, html: str) -> str:
        """Inject print rules into a complete proposal document."""
        return html.replace("</head>", f"{_PRINT_STYLE}</head>", 1) if "</head>" in html else _PRINT_STYLE + html

    async def render_pdf(self, html: str, base_url: str | Path | None = None) -> bytes:
        """Render PDF via WeasyPrint, falling back to installed Playwright Chromium."""
        async with _PDF_SEMAPHORE:
            prepared = self.render_html(html)
            try:
                from weasyprint import HTML

                return await asyncio.wait_for(
                    asyncio.to_thread(HTML(string=prepared, base_url=str(base_url or Path.cwd())).write_pdf),
                    timeout=30,
                )
            except Exception as weasy_error:
                try:
                    return await asyncio.wait_for(self._render_playwright(prepared), timeout=30)
                except Exception as fallback_error:
                    raise RuntimeError(
                        "PDF rendering unavailable: WeasyPrint and Playwright failed"
                    ) from fallback_error

    async def _render_playwright(self, html: str) -> bytes:
        try:
            from playwright.async_api import async_playwright
        except ImportError as exc:
            raise RuntimeError("PDF rendering requires WeasyPrint or Playwright") from exc
        async with async_playwright() as playwright:
            try:
                browser = await playwright.chromium.launch(headless=True)
            except Exception as exc:
                raise RuntimeError("PDF rendering unavailable: install WeasyPrint or Playwright Chromium") from exc
            try:
                page = await browser.new_page()
                await page.set_content(html, wait_until="load")
                return await page.pdf(format="A4", print_background=True)
            finally:
                await browser.close()

    def render_docx(self, html: str) -> bytes:
        """Render semantic text blocks when python-docx is installed."""
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX rendering is optional; install python-docx to enable it") from exc
        parser = _TextExtractor()
        parser.feed(html)
        document = Document()
        for tag, text in parser.blocks:
            if tag.startswith("h"):
                document.add_heading(text, level=min(int(tag[1]), 3))
            else:
                document.add_paragraph(text, style="List Bullet" if tag == "li" else None)
        output = BytesIO()
        document.save(output)
        return output.getvalue()


_renderer = PDFRenderer()


async def render_pdf(html: str, base_url: str | Path | None = None) -> bytes:
    """Render PDF through the shared bounded renderer."""
    return await _renderer.render_pdf(html, base_url)


def render_html(html: str) -> str:
    """Render printable HTML through the shared renderer."""
    return _renderer.render_html(html)


def render_docx(html: str) -> bytes:
    """Render optional DOCX through the shared renderer."""
    return _renderer.render_docx(html)
